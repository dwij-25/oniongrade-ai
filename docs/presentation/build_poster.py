import os
import math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

COLOR_LIGHT_BLUE = RGBColor(187, 223, 245)      # #BBDFF5 (Header & Footer)
COLOR_NAVY_TITLE = RGBColor(36, 39, 139)        # #24278B (Navy titles)
COLOR_BORDER_BLUE = RGBColor(93, 173, 226)      # #5DADE2 (Card border & accents)
COLOR_BORDER_TINT = RGBColor(180, 215, 240)     # Subtle card outline
COLOR_BOX_BG = RGBColor(255, 255, 255)          # Card background
COLOR_TEXT_BODY = RGBColor(45, 55, 72)          # #2D3748 Body text
COLOR_PLACEHOLDER = RGBColor(120, 140, 160)     # Placeholder text
COLOR_DIAGRAM_BG = RGBColor(234, 244, 252)      # Diagram step box fill
COLOR_DIAGRAM_BORDER = RGBColor(93, 173, 226)  # Diagram step box border
COLOR_DIAGRAM_TEXT = RGBColor(27, 42, 114)      # Diagram text
COLOR_LOGO_BG = RGBColor(248, 250, 252)         # Logo background

CORNER_RADIUS_INCHES = 0.50

def add_smooth_corner_bracket(slide, x, y, w, h, lh, lv, thick, color):
    r_out = Inches(CORNER_RADIUS_INCHES)
    thick_in = Inches(thick)
    r_in = r_out - thick_in
    cx = x + w - r_out
    cy = y + h - r_out

    pts = []
    # Bottom horizontal outer edge (starts at left of horizontal arm)
    pts.append((cx - lh + r_out, y + h))
    # Outer arc around bottom-right corner (from 90 deg at bottom to 0 deg at right)
    for deg in range(90, -5, -5):
        rad = math.radians(deg)
        pts.append((cx + r_out * math.cos(rad), cy + r_out * math.sin(rad)))
    # Right vertical outer edge (up to top of vertical arm)
    pts.append((x + w, cy - lv + r_out))
    # Right vertical inner edge
    pts.append((x + w - thick_in, cy - lv + r_out))
    # Inner arc around corner (from 0 deg back to 90 deg)
    for deg in range(0, 95, 5):
        rad = math.radians(deg)
        pts.append((cx + r_in * math.cos(rad), cy + r_in * math.sin(rad)))
    # Bottom horizontal inner edge
    pts.append((cx - lh + r_out, y + h - thick_in))

    builder = slide.shapes.build_freeform(pts[0][0], pts[0][1])
    builder.add_line_segments(pts[1:])
    bracket = builder.convert_to_shape()
    bracket.fill.solid()
    bracket.fill.fore_color.rgb = color
    bracket.line.fill.background()
    return bracket

def create_card_with_accents(slide, x, y, w, h, lh=Inches(3.8), lv=Inches(2.8)):
    # 1. Base Card with uniform corner radius
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    # Calculate adjustment to ensure exact 0.5 inch radius
    min_dim = min(w, h)
    card.adjustments[0] = CORNER_RADIUS_INCHES / (min_dim / 914400.0)
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_BOX_BG
    card.line.color.rgb = COLOR_BORDER_BLUE
    card.line.width = Pt(1.75)

    # 2. Left Accent Bar
    bar_w = Inches(0.18)
    bar_inset = Inches(0.06)
    r_margin = Inches(CORNER_RADIUS_INCHES + 0.15)
    bar_y = y + r_margin
    bar_h = h - (r_margin * 2)

    bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + bar_inset, bar_y, bar_w, bar_h)
    bar.adjustments[0] = 0.5
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_BORDER_BLUE
    bar.line.fill.background()

    # 3. Bottom-Right Smooth Corner Accent Tab
    add_smooth_corner_bracket(slide, x, y, w, h, lh=lh, lv=lv, thick=0.18, color=COLOR_BORDER_BLUE)

    return card

def build_poster(output_path):
    prs = Presentation()
    prs.slide_width = Inches(36.0)
    prs.slide_height = Inches(48.0)

    prs.core_properties.title = "OnionGrade AI: A Computer Vision and Deep Learning Framework for Real-Time Onion Quality Grading and AGMARK-Compliant Classification"
    prs.core_properties.subject = "National Seminar Academic Poster Presentation"
    prs.core_properties.author = "Gujarat Technological University"

    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # 1. HEADER BANNER
    header_h = Inches(1.65)
    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(36.0), header_h)
    header.fill.solid()
    header.fill.fore_color.rgb = COLOR_LIGHT_BLUE
    header.line.fill.background()

    tf_h = header.text_frame
    tf_h.word_wrap = True
    tf_h.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_h = tf_h.paragraphs[0]
    p_h.alignment = PP_ALIGN.CENTER
    r_h = p_h.add_run()
    run_text = "Gujarat Technological University – Ahmedabad"
    r_h.text = run_text
    r_h.font.name = "Times New Roman"
    r_h.font.size = Pt(36)
    r_h.font.bold = True
    r_h.font.color.rgb = COLOR_NAVY_TITLE

    # 2. INFO FIELDS BLOCK & LOGOS
    info_top = Inches(2.20)
    info_h = Inches(5.25)
    margin_x = Inches(1.50)

    # University Logo Placeholder (Top-Left, Rounded Square)
    logo_dim = Inches(5.25)
    uni_logo = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, margin_x, info_top, logo_dim, logo_dim)
    uni_logo.adjustments[0] = 0.12
    uni_logo.fill.solid()
    uni_logo.fill.fore_color.rgb = COLOR_LOGO_BG
    uni_logo.line.color.rgb = COLOR_BORDER_BLUE
    uni_logo.line.width = Pt(1.5)

    tf_uni = uni_logo.text_frame
    tf_uni.word_wrap = True
    tf_uni.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_u1 = tf_uni.paragraphs[0]
    p_u1.alignment = PP_ALIGN.CENTER
    ru1 = p_u1.add_run()
    ru1.text = "University Logo"
    ru1.font.name = "Times New Roman"
    ru1.font.size = Pt(22)
    ru1.font.bold = True
    ru1.font.color.rgb = COLOR_PLACEHOLDER

    p_u2 = tf_uni.add_paragraph()
    p_u2.alignment = PP_ALIGN.CENTER
    p_u2.space_before = Pt(6)
    ru2 = p_u2.add_run()
    ru2.text = "[Insert GTU Logo]"
    ru2.font.name = "Times New Roman"
    ru2.font.size = Pt(15)
    ru2.font.color.rgb = COLOR_PLACEHOLDER

    # Institute Logo Placeholder (Top-Right, Circle)
    inst_x = Inches(36.0) - margin_x - logo_dim
    inst_logo = slide.shapes.add_shape(MSO_SHAPE.OVAL, inst_x, info_top, logo_dim, logo_dim)
    inst_logo.fill.solid()
    inst_logo.fill.fore_color.rgb = COLOR_LOGO_BG
    inst_logo.line.color.rgb = COLOR_BORDER_BLUE
    inst_logo.line.width = Pt(1.5)

    tf_inst = inst_logo.text_frame
    tf_inst.word_wrap = True
    tf_inst.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_i1 = tf_inst.paragraphs[0]
    p_i1.alignment = PP_ALIGN.CENTER
    ri1 = p_i1.add_run()
    ri1.text = "Institute Logo"
    ri1.font.name = "Times New Roman"
    ri1.font.size = Pt(22)
    ri1.font.bold = True
    ri1.font.color.rgb = COLOR_PLACEHOLDER

    p_i2 = tf_inst.add_paragraph()
    p_i2.alignment = PP_ALIGN.CENTER
    p_i2.space_before = Pt(6)
    ri2 = p_i2.add_run()
    ri2.text = "[Insert College Logo]"
    ri2.font.name = "Times New Roman"
    ri2.font.size = Pt(15)
    ri2.font.color.rgb = COLOR_PLACEHOLDER

    # Four Stacked Info Fields (Centered)
    info_box_x = Inches(7.25)
    info_box_w = Inches(21.50)
    field_h = Inches(1.10)
    field_gap = Inches(0.28)

    info_data = [
        ("[Name of Students]", 28, False, COLOR_PLACEHOLDER),
        ("[Supervisor's and Co-supervisor's Name]", 28, False, COLOR_PLACEHOLDER),
        ("[Semester and Department]", 28, False, COLOR_PLACEHOLDER),
        ("Name of the Institute", 32, True, COLOR_NAVY_TITLE)
    ]

    for idx, (txt, fsize, bold_flag, fcolor) in enumerate(info_data):
        fy = info_top + idx * (field_h + field_gap)
        ibox = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, info_box_x, fy, info_box_w, field_h)
        ibox.adjustments[0] = 0.22
        ibox.fill.solid()
        ibox.fill.fore_color.rgb = COLOR_BOX_BG
        ibox.line.color.rgb = RGBColor(176, 212, 241)
        ibox.line.width = Pt(1.5)

        itf = ibox.text_frame
        itf.word_wrap = True
        itf.vertical_anchor = MSO_ANCHOR.MIDDLE
        ip = itf.paragraphs[0]
        ip.alignment = PP_ALIGN.CENTER
        ir = ip.add_run()
        ir.text = txt
        ir.font.name = "Times New Roman"
        ir.font.size = Pt(fsize)
        ir.font.bold = bold_flag
        ir.font.color.rgb = fcolor

    # 3. CONTENT GRID (2 COLUMNS x 3 ROWS)
    grid_top = Inches(8.05)
    row_h = Inches(10.30)
    row_gap = Inches(0.50)
    col_w = Inches(15.90)
    col_gap = Inches(1.20)
    col1_x = margin_x
    col2_x = margin_x + col_w + col_gap

    # ROW 1 LEFT: Introduction
    create_card_with_accents(slide, col1_x, grid_top, col_w, row_h)
    tx_intro = slide.shapes.add_textbox(col1_x + Inches(0.65), grid_top + Inches(0.45), col_w - Inches(1.30), row_h - Inches(0.90))
    tf_intro = tx_intro.text_frame
    tf_intro.word_wrap = True
    p_tit = tf_intro.paragraphs[0]
    p_tit.space_after = Pt(16)
    r_tit = p_tit.add_run()
    r_tit.text = "Introduction:"
    r_tit.font.name = "Times New Roman"
    r_tit.font.size = Pt(22)
    r_tit.font.bold = True
    r_tit.font.color.rgb = COLOR_NAVY_TITLE

    p_body = tf_intro.add_paragraph()
    p_body.line_spacing = 1.35
    r_body = p_body.add_run()
    r_body.text = (
        "Onion quality grading at government procurement centres is currently performed "
        "manually, leading to inconsistent Grade A / URS classification and recurring "
        "farmer–government disputes. This project proposes OnionGrade AI, a mobile "
        "computer-vision system that photographs a lot of onions, automatically detects "
        "damage, rot, sprouting and undersizing, and instantly estimates the percentage of "
        "Grade A and URS onions in the lot — replacing subjective visual inspection with a "
        "repeatable, rule-based AI assessment that reduces human bias and improves transparency."
    )
    r_body.font.name = "Times New Roman"
    r_body.font.size = Pt(18)
    r_body.font.color.rgb = COLOR_TEXT_BODY

    # ROW 1 RIGHT: Block Diagram / Implementation Details
    create_card_with_accents(slide, col2_x, grid_top, col_w, row_h)
    tx_diag_title = slide.shapes.add_textbox(col2_x + Inches(0.65), grid_top + Inches(0.45), col_w - Inches(1.30), Inches(0.80))
    tf_diag_t = tx_diag_title.text_frame
    tf_diag_t.word_wrap = True
    p_diag_t = tf_diag_t.paragraphs[0]
    r_diag_t = p_diag_t.add_run()
    r_diag_t.text = "Block Diagram / Implementation Details:"
    r_diag_t.font.name = "Times New Roman"
    r_diag_t.font.size = Pt(22)
    r_diag_t.font.bold = True
    r_diag_t.font.color.rgb = COLOR_NAVY_TITLE

    # 5 Connected Diagram Boxes with Downward Arrows
    diag_box_w = Inches(12.80)
    diag_box_h = Inches(1.15)
    diag_box_x = col2_x + (col_w - diag_box_w) / 2
    diag_start_y = grid_top + Inches(1.45)
    arrow_gap = Inches(0.55)
    arrow_w = Inches(0.45)
    arrow_h = Inches(0.38)

    diag_steps = [
        "Capture (smartphone)",
        "Segmentation (per-bulb crop)",
        "Feature extraction",
        "Rule engine (AGMARK table)",
        "Instant report (% A / URS / Reject)"
    ]

    for idx, step_text in enumerate(diag_steps):
        box_y = diag_start_y + idx * (diag_box_h + arrow_gap)
        dbox = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, diag_box_x, box_y, diag_box_w, diag_box_h)
        dbox.adjustments[0] = 0.20
        dbox.fill.solid()
        dbox.fill.fore_color.rgb = COLOR_DIAGRAM_BG
        dbox.line.color.rgb = COLOR_DIAGRAM_BORDER
        dbox.line.width = Pt(1.75)

        dtf = dbox.text_frame
        dtf.word_wrap = True
        dtf.vertical_anchor = MSO_ANCHOR.MIDDLE
        dp = dtf.paragraphs[0]
        dp.alignment = PP_ALIGN.CENTER
        dr = dp.add_run()
        dr.text = step_text
        dr.font.name = "Times New Roman"
        dr.font.size = Pt(18)
        dr.font.bold = True
        dr.font.color.rgb = COLOR_DIAGRAM_TEXT

        if idx < len(diag_steps) - 1:
            arr_y = box_y + diag_box_h + (arrow_gap - arrow_h) / 2
            arr_x = col2_x + (col_w - arrow_w) / 2
            arrow = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, arr_x, arr_y, arrow_w, arrow_h)
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = COLOR_BORDER_BLUE
            arrow.line.fill.background()

    # ROW 2 LEFT: Literature Review / Market Survey
    row2_y = grid_top + row_h + row_gap
    create_card_with_accents(slide, col1_x, row2_y, col_w, row_h)
    tx_lit = slide.shapes.add_textbox(col1_x + Inches(0.65), row2_y + Inches(0.45), col_w - Inches(1.30), row_h - Inches(0.90))
    tf_lit = tx_lit.text_frame
    tf_lit.word_wrap = True
    p_lit_t = tf_lit.paragraphs[0]
    p_lit_t.space_after = Pt(16)
    r_lit_t = p_lit_t.add_run()
    r_lit_t.text = "Literature Review / Market Survey:"
    r_lit_t.font.name = "Times New Roman"
    r_lit_t.font.size = Pt(22)
    r_lit_t.font.bold = True
    r_lit_t.font.color.rgb = COLOR_NAVY_TITLE

    p_lit_b = tf_lit.add_paragraph()
    p_lit_b.line_spacing = 1.35
    r_lit_b = p_lit_b.add_run()
    r_lit_b.text = (
        "A DNN classifying onions on geometric and skin-damage features achieved 91.85% "
        "accuracy (Acta Sci. Pol. Hortorum Cultus). MobileNetV2 transfer learning is validated "
        "for low-cost, on-device onion classification; DenseNet121 outperforms other CNN "
        "backbones for Indian onion variety ID (ICAR-DOGR study). Feature-adaptive anomaly "
        "detection enables identification of unseen defect types. AGMARK defines Grade A and "
        "the Under Relaxed Specification (URS) tier used in NAFED/NCCF procurement — 2026 "
        "reports confirm active farmer disputes over inconsistent grading."
    )
    r_lit_b.font.name = "Times New Roman"
    r_lit_b.font.size = Pt(18)
    r_lit_b.font.color.rgb = COLOR_TEXT_BODY

    # ROW 2 RIGHT: Test Set-up and Results
    create_card_with_accents(slide, col2_x, row2_y, col_w, row_h)
    tx_test = slide.shapes.add_textbox(col2_x + Inches(0.65), row2_y + Inches(0.45), col_w - Inches(1.30), row_h - Inches(0.90))
    tf_test = tx_test.text_frame
    tf_test.word_wrap = True
    p_test_t = tf_test.paragraphs[0]
    p_test_t.space_after = Pt(16)
    r_test_t = p_test_t.add_run()
    r_test_t.text = "Test Set-up and Results:"
    r_test_t.font.name = "Times New Roman"
    r_test_t.font.size = Pt(22)
    r_test_t.font.bold = True
    r_test_t.font.color.rgb = COLOR_NAVY_TITLE

    p_test_b = tf_test.add_paragraph()
    p_test_b.line_spacing = 1.35
    r_test_b = p_test_b.add_run()
    r_test_b.text = (
        "Labelled onion images collected under varied lighting/background, benchmarked "
        "against inspector-assigned ground truth on a mid-range Android smartphone (₹8–10k "
        "class). Backbone: MobileNetV2 (on-device) + DenseNet121 (cloud verifier). Based on "
        "comparable published studies, expected performance: 90%+ accuracy, sub-2-second "
        "per-lot inference. Agreement with human inspectors measured via Cohen's kappa; "
        "false-accept and false-reject rates tracked separately."
    )
    r_test_b.font.name = "Times New Roman"
    r_test_b.font.size = Pt(18)
    r_test_b.font.color.rgb = COLOR_TEXT_BODY

    # ROW 3 LEFT: Methodology
    row3_y = row2_y + row_h + row_gap
    create_card_with_accents(slide, col1_x, row3_y, col_w, row_h)
    tx_meth = slide.shapes.add_textbox(col1_x + Inches(0.65), row3_y + Inches(0.45), col_w - Inches(1.30), row_h - Inches(0.90))
    tf_meth = tx_meth.text_frame
    tf_meth.word_wrap = True
    p_meth_t = tf_meth.paragraphs[0]
    p_meth_t.space_after = Pt(16)
    r_meth_t = p_meth_t.add_run()
    r_meth_t.text = "Methodology:"
    r_meth_t.font.name = "Times New Roman"
    r_meth_t.font.size = Pt(22)
    r_meth_t.font.bold = True
    r_meth_t.font.color.rgb = COLOR_NAVY_TITLE

    p_meth_b = tf_meth.add_paragraph()
    p_meth_b.line_spacing = 1.35
    r_meth_b = p_meth_b.add_run()
    r_meth_b.text = (
        "Pipeline: (1) Capture — smartphone photo of onion tray; (2) Segmentation — per-bulb "
        "contour/blob detection; (3) Feature extraction — diameter, skin gloss, dark-patch "
        "(rot) ratio, sprout presence, shape irregularity; (4) Rule-based classification — "
        "features checked against an editable, AGMARK-aligned threshold table (Grade A / URS "
        "/ Reject); (5) Report generation — instant lot-level percentages plus annotated "
        "photo showing why each bulb was downgraded. Ground truth collected with ICAR-DOGR "
        "and pilot APMC centres."
    )
    r_meth_b.font.name = "Times New Roman"
    r_meth_b.font.size = Pt(18)
    r_meth_b.font.color.rgb = COLOR_TEXT_BODY

    # ROW 3 RIGHT: Conclusions / Summary and Future Work with Acknowledgment
    create_card_with_accents(slide, col2_x, row3_y, col_w, row_h)
    tx_conc = slide.shapes.add_textbox(col2_x + Inches(0.65), row3_y + Inches(0.45), col_w - Inches(1.30), row_h - Inches(0.90))
    tf_conc = tx_conc.text_frame
    tf_conc.word_wrap = True
    p_conc_t = tf_conc.paragraphs[0]
    p_conc_t.space_after = Pt(16)
    r_conc_t = p_conc_t.add_run()
    r_conc_t.text = "Conclusions / Summary and Future Work with Acknowledgment:"
    r_conc_t.font.name = "Times New Roman"
    r_conc_t.font.size = Pt(22)
    r_conc_t.font.bold = True
    r_conc_t.font.color.rgb = COLOR_NAVY_TITLE

    p_conc_b = tf_conc.add_paragraph()
    p_conc_b.line_spacing = 1.35
    r_conc_b = p_conc_b.add_run()
    r_conc_b.text = (
        "OnionGrade AI shows a smartphone-based computer vision pipeline can replace "
        "subjective manual onion grading with a fast, explainable, auditable process — "
        "directly addressing recent NAFED/NCCF procurement disputes. Future work: hybrid "
        "edge-cloud architecture with an immutable audit trail, and federated learning so the "
        "model adapts to regional onion varieties without centralising farmer data. "
        "Acknowledgment: ICAR-Directorate of Onion & Garlic Research (ICAR-DOGR) and partner "
        "APMC centres for domain guidance."
    )
    r_conc_b.font.name = "Times New Roman"
    r_conc_b.font.size = Pt(18)
    r_conc_b.font.color.rgb = COLOR_TEXT_BODY

    # 4. FULL-WIDTH REFERENCES BOX
    ref_y = row3_y + row_h + Inches(0.55)
    ref_w = Inches(33.0)
    ref_h = Inches(5.20)
    create_card_with_accents(slide, margin_x, ref_y, ref_w, ref_h, lh=Inches(5.0), lv=Inches(2.0))

    # Reference box Title
    tx_ref_t = slide.shapes.add_textbox(margin_x + Inches(0.65), ref_y + Inches(0.35), ref_w - Inches(1.30), Inches(0.60))
    tf_ref_t = tx_ref_t.text_frame
    tf_ref_t.word_wrap = True
    p_ref_t = tf_ref_t.paragraphs[0]
    r_ref_t = p_ref_t.add_run()
    r_ref_t.text = "References:"
    r_ref_t.font.name = "Times New Roman"
    r_ref_t.font.size = Pt(22)
    r_ref_t.font.bold = True
    r_ref_t.font.color.rgb = COLOR_NAVY_TITLE

    # Two columns for 6 references
    ref_col_w = Inches(15.40)
    ref_col_gap = Inches(0.80)
    ref_col1_x = margin_x + Inches(0.65)
    ref_col2_x = ref_col1_x + ref_col_w + ref_col_gap
    ref_body_y = ref_y + Inches(1.05)
    ref_body_h = ref_h - Inches(1.20)

    refs = [
        ("1.", "Author(s), \"Development and application of a model for automatic evaluation and classification of onions using a DNN,\" Acta Sci. Pol. Hortorum Cultus, pp. xx-xx, 2023."),
        ("2.", "Author(s), \"Image-based identification of onion varieties using deep learning techniques,\" Vegetable Science, pp. xx-xx, 2024."),
        ("3.", "Author(s), \"Onion Image Classification Based on Transfer Learning with MobileNetV2,\" ResearchGate, 2026."),
        ("4.", "Author(s), \"Feature-adaptive anomaly detection model for onion inspection system,\" ScienceDirect, pp. xx-xx, 2025."),
        ("5.", "Directorate of Marketing & Inspection, \"AGMARK Grades and Standards,\" Govt. of India, 2019."),
        ("6.", "Free Press Journal, \"Nashik: Farmers Allege Quality Onion Rejections in NAFED Procurement, Demand Transparent Grading System,\" June 2026.")
    ]

    # Col 1: Refs 1-3
    tx_ref1 = slide.shapes.add_textbox(ref_col1_x, ref_body_y, ref_col_w, ref_body_h)
    tf_ref1 = tx_ref1.text_frame
    tf_ref1.word_wrap = True
    for i in range(3):
        num, rtext = refs[i]
        p = tf_ref1.paragraphs[0] if i == 0 else tf_ref1.add_paragraph()
        p.space_after = Pt(12)
        p.line_spacing = 1.25
        r_num = p.add_run()
        r_num.text = f"{num} "
        r_num.font.name = "Times New Roman"
        r_num.font.size = Pt(14.5)
        r_num.font.bold = True
        r_num.font.color.rgb = COLOR_NAVY_TITLE

        r_body = p.add_run()
        r_body.text = rtext
        r_body.font.name = "Times New Roman"
        r_body.font.size = Pt(14.5)
        r_body.font.color.rgb = COLOR_TEXT_BODY

    # Col 2: Refs 4-6
    tx_ref2 = slide.shapes.add_textbox(ref_col2_x, ref_body_y, ref_col_w, ref_body_h)
    tf_ref2 = tx_ref2.text_frame
    tf_ref2.word_wrap = True
    for i in range(3, 6):
        num, rtext = refs[i]
        p = tf_ref2.paragraphs[0] if i == 3 else tf_ref2.add_paragraph()
        p.space_after = Pt(12)
        p.line_spacing = 1.25
        r_num = p.add_run()
        r_num.text = f"{num} "
        r_num.font.name = "Times New Roman"
        r_num.font.size = Pt(14.5)
        r_num.font.bold = True
        r_num.font.color.rgb = COLOR_NAVY_TITLE

        r_body = p.add_run()
        r_body.text = rtext
        r_body.font.name = "Times New Roman"
        r_body.font.size = Pt(14.5)
        r_body.font.color.rgb = COLOR_TEXT_BODY

    # 5. FOOTER BANNER
    footer_top = Inches(46.35)
    footer_h = Inches(1.65)
    footer = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), footer_top, Inches(36.0), footer_h)
    footer.fill.solid()
    footer.fill.fore_color.rgb = COLOR_LIGHT_BLUE
    footer.line.fill.background()

    tf_footer = footer.text_frame
    tf_footer.word_wrap = True
    tf_footer.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_footer = tf_footer.paragraphs[0]
    p_footer.alignment = PP_ALIGN.CENTER
    r_footer = p_footer.add_run()
    r_footer.text = "AI-Powered Research & Innovation"
    r_footer.font.name = "Times New Roman"
    r_footer.font.size = Pt(18)
    r_footer.font.bold = True
    r_footer.font.color.rgb = COLOR_NAVY_TITLE

    prs.save(output_path)
    print(f"Poster rebuilt successfully: {output_path}")

if __name__ == "__main__":
    out_file = os.path.abspath("OnionGrade_AI_Academic_Poster.pptx")
    build_poster(out_file)
